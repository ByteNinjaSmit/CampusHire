"""Admin module tests: RBAC, health, metrics, audit logs, compliance, export / import jobs."""

import csv
import io
import uuid
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import redis as redis_core
from app.core.middleware import metrics_key
from app.modules.admin import exporter, health, importer, service
from app.modules.admin import schemas as s
from app.modules.admin.models import AuditLog, Job, LoginEvent, PolicyViolation
from app.modules.companies.models import Company
from app.modules.students.models import Student
from app.modules.users.models import User
from tests.factories import (
    auth_headers,
    make_admin,
    make_application,
    make_company,
    make_faculty,
    make_internship,
    make_resume,
    make_student,
    make_user,
)

API = "/api/v1/admin"


@pytest.fixture
def session_factory(db_conn):
    """Sessions for run_*_job on the test connection (same transaction as the API requests)."""

    def factory():
        return AsyncSession(bind=db_conn, join_transaction_mode="create_savepoint", expire_on_commit=False)

    return factory


@pytest.fixture
async def admin(db):
    return await make_admin(db)


def xlsx_bytes(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_bytes(rows: list[list[str]]) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode()


# ---- RBAC ----------------------------------------------------------------------------------------------------------
ADMIN_ROUTES = [
    ("GET", "/health"),
    ("GET", "/metrics"),
    ("GET", "/audit-logs"),
    ("GET", "/compliance/policies"),
    ("GET", "/compliance/violations"),
    ("GET", "/compliance/documents"),
    ("POST", "/compliance/scan"),
    ("POST", "/export"),
    ("POST", "/import"),
    ("GET", "/import/template/students"),
    ("GET", "/jobs"),
]


@pytest.mark.parametrize("role", ["STUDENT", "FACULTY", "COMPANY"])
@pytest.mark.parametrize("method,path", ADMIN_ROUTES)
async def test_non_admin_forbidden(client, db, role, method, path):
    user = await make_user(db, role)
    resp = await client.request(method, API + path, headers=auth_headers(user))
    assert resp.status_code == 403


@pytest.mark.parametrize("method,path", ADMIN_ROUTES)
async def test_anonymous_unauthorised(client, method, path):
    assert (await client.request(method, API + path)).status_code == 401


# ---- health ---------------------------------------------------------------------------------------------------------
async def test_health_reports_components(client, db, admin, monkeypatch):
    async def fake_minio():
        return s.MinioHealth(
            status="ok",
            latency_ms=1.0,
            buckets=[s.BucketHealth(bucket="resumes", status="ok", object_count=2, size_bytes=300)],
            total_size_bytes=300,
        )

    async def fake_celery():
        return s.CeleryHealth(status="down", detail="no workers responded")

    async def fake_depth():
        return 4

    monkeypatch.setattr(health, "check_minio", fake_minio)
    monkeypatch.setattr(health, "check_celery", fake_celery)
    monkeypatch.setattr(health, "queue_depth", fake_depth)
    resp = await client.get(f"{API}/health", headers=auth_headers(admin))
    assert resp.status_code == 200
    body = resp.json()
    assert body["db"]["status"] == "ok" and body["db"]["latency_ms"] is not None
    assert body["redis"]["status"] == "ok"
    assert body["minio"]["total_size_bytes"] == 300 and body["minio"]["buckets"][0]["bucket"] == "resumes"
    assert body["celery"]["status"] == "down"
    assert body["queue_depth"] == 4
    assert body["status"] == "degraded"


async def test_health_down_when_redis_down(client, db, admin, monkeypatch):
    async def bad_redis():
        return s.ComponentHealth(status="down", detail="boom")

    async def ok_minio():
        return s.MinioHealth(status="ok")

    async def ok_celery():
        return s.CeleryHealth(status="ok", workers=["w1"])

    async def depth():
        return None

    monkeypatch.setattr(health, "check_redis", bad_redis)
    monkeypatch.setattr(health, "check_minio", ok_minio)
    monkeypatch.setattr(health, "check_celery", ok_celery)
    monkeypatch.setattr(health, "queue_depth", depth)
    body = (await client.get(f"{API}/health", headers=auth_headers(admin))).json()
    assert body["status"] == "down"
    assert body["queue_depth"] is None


# ---- metrics ----------------------------------------------------------------------------------------------------------
async def test_metrics_from_redis_and_login_events(client, db, admin):
    r = redis_core.get_redis()
    key = metrics_key(datetime.now(UTC))
    await r.hset(
        key,
        mapping={"count": 10, "errors": 1, "lat_le_5": 4, "lat_le_10": 4, "lat_le_100": 2, "lat_sum_ms": 120.0},
    )
    db.add_all(
        [
            LoginEvent(user_id=admin.id, email=admin.email, success=True),
            LoginEvent(user_id=admin.id, email=admin.email, success=True),
            LoginEvent(user_id=None, email="x@y.com", success=False),
        ]
    )
    await db.flush()
    resp = await client.get(f"{API}/metrics?minutes=5&login_days=3", headers=auth_headers(admin))
    assert resp.status_code == 200
    body = resp.json()
    assert body["minutes"] == 5 and len(body["points"]) == 5
    assert body["totals"]["requests"] >= 10 and body["totals"]["errors"] >= 1
    assert body["totals"]["p50_ms"] is not None and body["totals"]["p95_ms"] > body["totals"]["p50_ms"]
    assert len(body["login_trend"]) == 3
    assert body["login_trend"][-1]["success"] == 2 and body["login_trend"][-1]["failed"] == 1
    assert body["active_users_24h"] == 1 and body["logins_last_hour"] == 2


async def test_metrics_minutes_bounds(client, admin):
    assert (await client.get(f"{API}/metrics?minutes=0", headers=auth_headers(admin))).status_code == 422
    assert (await client.get(f"{API}/metrics?minutes=500", headers=auth_headers(admin))).status_code == 422


# ---- audit logs ---------------------------------------------------------------------------------------------------------
async def test_audit_log_filters(client, db, admin):
    other = await make_admin(db)
    db.add_all(
        [
            AuditLog(actor_id=admin.id, action="user.create", entity_type="user", entity_id=uuid.uuid4()),
            AuditLog(actor_id=admin.id, action="internship.approve", entity_type="internship"),
            AuditLog(actor_id=other.id, action="user.deactivate", entity_type="user"),
        ]
    )
    await db.flush()
    h = auth_headers(admin)
    resp = await client.get(f"{API}/audit-logs?actor_id={other.id}", headers=h)
    assert [i["action"] for i in resp.json()["items"]] == ["user.deactivate"]
    assert resp.json()["items"][0]["actor"]["id"] == str(other.id)
    resp = await client.get(f"{API}/audit-logs?entity_type=user", headers=h)
    assert resp.json()["total"] == 2
    resp = await client.get(f"{API}/audit-logs?action=user.", headers=h)
    assert {i["action"] for i in resp.json()["items"]} == {"user.create", "user.deactivate"}
    future = (datetime.now(UTC) + timedelta(days=1)).isoformat()
    resp = await client.get(f"{API}/audit-logs", headers=h, params={"from": future})
    assert resp.json()["total"] == 0
    past = (datetime.now(UTC) - timedelta(days=1)).isoformat()
    resp = await client.get(f"{API}/audit-logs", headers=h, params={"from": past, "page_size": 2})
    body = resp.json()
    assert body["page_size"] == 2 and len(body["items"]) == 2 and body["pages"] == 2


# ---- compliance policies ----------------------------------------------------------------------------------------------------
async def test_policies_list_toggle_create(client, db, admin):
    h = auth_headers(admin)
    resp = await client.get(f"{API}/compliance/policies", headers=h)
    assert resp.status_code == 200
    policies = {p["code"]: p for p in resp.json()}
    assert {"RESUME_UNVERIFIED_7D", "INTERNSHIP_PAST_DEADLINE_OPEN", "UNPAID_LONG_INTERNSHIP"} <= set(policies)
    pid = policies["UNPAID_LONG_INTERNSHIP"]["id"]
    resp = await client.patch(f"{API}/compliance/policies/{pid}", headers=h, json={"is_active": False})
    assert resp.status_code == 200 and resp.json()["is_active"] is False
    resp = await client.post(
        f"{API}/compliance/policies", headers=h, json={"code": "custom_rule", "name": "Custom rule", "severity": "LOW"}
    )
    assert resp.status_code == 201 and resp.json()["code"] == "CUSTOM_RULE"
    dup = await client.post(f"{API}/compliance/policies", headers=h, json={"code": "CUSTOM_RULE", "name": "Again"})
    assert dup.status_code == 409
    bad = await client.post(f"{API}/compliance/policies", headers=h, json={"code": "x", "name": "N"})
    assert bad.status_code == 422
    missing = await client.patch(f"{API}/compliance/policies/{uuid.uuid4()}", headers=h, json={"is_active": True})
    assert missing.status_code == 404
    assert (
        await db.execute(select(func.count()).select_from(AuditLog).where(AuditLog.action.like("compliance.policy.%")))
    ).scalar_one() == 2


# ---- compliance scan / violations ---------------------------------------------------------------------------------------------
async def _seed_violating_data(db):
    faculty = await make_faculty(db)
    student = await make_student(db)
    internship = await make_internship(db, faculty, stipend_monthly=0, weeks=16)  # unpaid + long
    resume = await make_resume(db, student, verification_status="PENDING")
    resume.created_at = datetime.now(UTC) - timedelta(days=10)
    await make_application(db, student, internship, resume=resume)
    expired = await make_internship(db, faculty, deadline_in_days=-2)  # APPROVED but deadline passed
    await db.flush()
    return resume, internship, expired


async def test_scan_creates_violations_and_is_idempotent(client, db, admin):
    resume, unpaid, expired = await _seed_violating_data(db)
    first = await service.run_compliance_scan(db)
    assert first["by_policy"]["RESUME_UNVERIFIED_7D"] == 1
    assert first["by_policy"]["UNPAID_LONG_INTERNSHIP"] == 1
    assert first["by_policy"]["INTERNSHIP_PAST_DEADLINE_OPEN"] == 1
    assert first["new_violations"] >= 3
    total = (await db.execute(select(func.count()).select_from(PolicyViolation))).scalar_one()
    second = await service.run_compliance_scan(db)
    assert second["new_violations"] == 0 and second["auto_resolved"] == 0
    assert (await db.execute(select(func.count()).select_from(PolicyViolation))).scalar_one() == total

    h = auth_headers(admin)
    resp = await client.get(f"{API}/compliance/violations?status=OPEN&policy_code=RESUME_UNVERIFIED_7D", headers=h)
    assert resp.status_code == 200
    items = resp.json()["items"]
    assert len(items) == 1 and items[0]["entity_id"] == str(resume.id)
    assert items[0]["policy"]["code"] == "RESUME_UNVERIFIED_7D" and items[0]["details"]["days_pending"] >= 10


async def test_violation_resolve_dismiss_and_auto_resolve(client, db, admin):
    resume, unpaid, expired = await _seed_violating_data(db)
    await service.run_compliance_scan(db)
    h = auth_headers(admin)
    items = (await client.get(f"{API}/compliance/violations?policy_code=UNPAID_LONG_INTERNSHIP", headers=h)).json()[
        "items"
    ]
    vid = items[0]["id"]
    resp = await client.patch(
        f"{API}/compliance/violations/{vid}", headers=h, json={"status": "RESOLVED", "note": "ok"}
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["status"] == "RESOLVED" and body["note"] == "ok" and body["resolved_by"]["id"] == str(admin.id)
    resp = await client.patch(f"{API}/compliance/violations/{vid}", headers=h, json={"status": "OPEN"})
    assert resp.json()["status"] == "OPEN" and resp.json()["resolved_by"] is None
    resp = await client.patch(
        f"{API}/compliance/violations/{vid}", headers=h, json={"status": "DISMISSED", "note": "n/a"}
    )
    assert resp.json()["status"] == "DISMISSED"
    assert (
        await client.patch(f"{API}/compliance/violations/{uuid.uuid4()}", headers=h, json={"status": "OPEN"})
    ).status_code == 404
    assert (
        await client.patch(f"{API}/compliance/violations/{vid}", headers=h, json={"status": "BOGUS"})
    ).status_code == 422

    # a dismissed violation is not reopened by the next scan; fixing the cause auto-resolves open ones
    expired.status = "CLOSED"
    await db.flush()
    again = await service.run_compliance_scan(db)
    assert again["auto_resolved"] == 1
    open_ids = (await client.get(f"{API}/compliance/violations?status=OPEN", headers=h)).json()["items"]
    assert str(expired.id) not in {i["entity_id"] for i in open_ids}
    dismissed = (await client.get(f"{API}/compliance/violations?status=DISMISSED", headers=h)).json()
    assert dismissed["total"] == 1


async def test_inactive_policy_is_skipped(client, db, admin):
    await _seed_violating_data(db)
    h = auth_headers(admin)
    policies = (await client.get(f"{API}/compliance/policies", headers=h)).json()
    pid = next(p["id"] for p in policies if p["code"] == "UNPAID_LONG_INTERNSHIP")
    await client.patch(f"{API}/compliance/policies/{pid}", headers=h, json={"is_active": False})
    result = await service.run_compliance_scan(db)
    assert "UNPAID_LONG_INTERNSHIP" not in result["by_policy"]


async def test_scan_endpoint_queues_task_and_job_runs(client, db, admin, captured, session_factory):
    await _seed_violating_data(db)
    resp = await client.post(f"{API}/compliance/scan", headers=auth_headers(admin))
    assert resp.status_code == 202
    job = resp.json()
    assert job["type"] == "COMPLIANCE_SCAN" and job["status"] == "QUEUED"
    assert ("maintenance.compliance_scan", {"job_id": job["id"]}) in captured.tasks
    result = await service.run_compliance_scan_job(job["id"], session_factory=session_factory)
    assert result and result["new_violations"] >= 3
    got = (await client.get(f"{API}/jobs/{job['id']}", headers=auth_headers(admin))).json()
    assert got["status"] == "SUCCEEDED" and got["progress"] == 100 and got["result"]["open_total"] >= 3
    # scheduled run without a job id creates its own job row
    assert await service.run_compliance_scan_job(None, session_factory=session_factory) is not None
    jobs = (await client.get(f"{API}/jobs?type=COMPLIANCE_SCAN", headers=auth_headers(admin))).json()
    assert jobs["total"] == 2


async def test_document_verification_queue(client, db, admin):
    student = await make_student(db)
    pending = await make_resume(db, student, verification_status="PENDING", filename="a.pdf")
    await make_resume(db, student, verification_status="VERIFIED", filename="b.pdf")
    await make_resume(db, student, status="PENDING_UPLOAD", filename="c.pdf")
    h = auth_headers(admin)
    body = (await client.get(f"{API}/compliance/documents", headers=h)).json()
    assert [d["id"] for d in body["items"]] == [str(pending.id)]
    assert body["items"][0]["owner"]["email"] == student.email
    verified = (await client.get(f"{API}/compliance/documents?verification_status=VERIFIED", headers=h)).json()
    assert verified["total"] == 1


# ---- export -------------------------------------------------------------------------------------------------------------------------
async def test_export_request_validation(client, admin):
    h = auth_headers(admin)
    assert (
        await client.post(f"{API}/export", headers=h, json={"entity": "payroll", "format": "csv"})
    ).status_code == 422
    assert (
        await client.post(f"{API}/export", headers=h, json={"entity": "students", "format": "pdf"})
    ).status_code == 422
    bad = await client.post(
        f"{API}/export", headers=h, json={"entity": "students", "format": "csv", "filters": {"nope": 1}}
    )
    assert bad.status_code == 422 and "filters.nope" in str(bad.json())
    bad = await client.post(
        f"{API}/export", headers=h, json={"entity": "applications", "format": "csv", "filters": {"status": "hired"}}
    )
    assert bad.status_code == 422


async def test_export_students_csv_job(client, db, admin, captured, fake_storage, session_factory):
    await make_student(db, department="CSE", gpa="3.9", full_name="=cmd|' /C calc'!A0")
    await make_student(db, department="ECE", gpa="2.5")
    h = auth_headers(admin)
    resp = await client.post(
        f"{API}/export", headers=h, json={"entity": "students", "format": "csv", "filters": {"department": "CSE"}}
    )
    assert resp.status_code == 202
    job = resp.json()
    assert job["type"] == "DATA_EXPORT" and job["status"] == "QUEUED" and job["download_url"] is None
    assert ("data.export", {"job_id": job["id"]}) in captured.tasks

    await exporter.run_export_job(job["id"], session_factory=session_factory, store=fake_storage)
    done = (await client.get(f"{API}/jobs/{job['id']}", headers=h)).json()
    assert done["status"] == "SUCCEEDED" and done["progress"] == 100 and done["finished_at"]
    assert done["result"]["rows_ok"] == 1 and done["download_url"].startswith("http://fake-storage.local/reports/")
    assert any(m[1]["type"] == "job" for m in captured.messages)

    key = next(k for k in fake_storage.objects if k.endswith(".csv"))
    rows = list(csv.reader(io.StringIO(fake_storage.objects[key]["data"].decode("utf-8-sig"))))
    assert rows[0][:5] == ["user_id", "email", "full_name", "phone", "department"]
    assert len(rows) == 2
    assert rows[1][2].startswith("'=")  # formula injection neutralised
    dl = await client.get(f"{API}/jobs/{job['id']}/download", headers=h)
    assert dl.status_code == 200 and dl.json()["expires_in"] == 300
    # export job row created a documents row owned by the admin
    doc = await db.get(Job, uuid.UUID(job["id"]))
    assert doc is not None and doc.result_document_id is not None


@pytest.mark.parametrize("entity", ["students", "companies", "internships", "applications", "feedback"])
async def test_export_every_entity_xlsx(client, db, admin, fake_storage, session_factory, entity):
    faculty = await make_faculty(db)
    student = await make_student(db)
    internship = await make_internship(db, faculty)
    await make_application(db, student, internship)
    h = auth_headers(admin)
    job = (await client.post(f"{API}/export", headers=h, json={"entity": entity, "format": "xlsx"})).json()
    await exporter.run_export_job(job["id"], session_factory=session_factory, store=fake_storage)
    done = (await client.get(f"{API}/jobs/{job['id']}", headers=h)).json()
    assert done["status"] == "SUCCEEDED", done
    key = next(k for k in fake_storage.objects if k.endswith(".xlsx"))
    wb = load_workbook(io.BytesIO(fake_storage.objects[key]["data"]))
    ws = wb.active
    assert ws.max_row == (1 if entity == "feedback" else ws.max_row) and ws.max_row >= 1
    if entity != "feedback":
        assert ws.max_row >= 2


async def test_export_job_failure_is_recorded(client, db, admin, fake_storage, session_factory):
    job = Job(id=uuid.uuid4(), type="DATA_EXPORT", status="QUEUED", requested_by=admin.id, params={"entity": "bogus"})
    db.add(job)
    await db.flush()
    await exporter.run_export_job(job.id, session_factory=session_factory, store=fake_storage)
    got = (await client.get(f"{API}/jobs/{job.id}", headers=auth_headers(admin))).json()
    assert got["status"] == "FAILED" and "Invalid export parameters" in got["error"]


# ---- import ---------------------------------------------------------------------------------------------------------------------
async def test_import_template_download(client, admin):
    resp = await client.get(f"{API}/import/template/students", headers=auth_headers(admin))
    assert resp.status_code == 200 and "spreadsheetml" in resp.headers["content-type"]
    ws = load_workbook(io.BytesIO(resp.content)).worksheets[0]
    assert [c.value for c in ws[1]][:4] == ["email", "full_name", "phone", "department"]
    assert (await client.get(f"{API}/import/template/companies", headers=auth_headers(admin))).status_code == 200
    assert (await client.get(f"{API}/import/template/internships", headers=auth_headers(admin))).status_code == 422


STUDENT_HEADER = ["email", "full_name", "phone", "department", "gpa", "enrollment_no", "graduation_year"]


async def _post_import(client, admin, entity, filename, content, ctype="text/csv"):
    return await client.post(
        f"{API}/import",
        headers=auth_headers(admin),
        data={"entity": entity},
        files={"file": (filename, content, ctype)},
    )


async def test_import_students_csv_with_bad_rows(client, db, admin, captured, fake_storage, session_factory):
    existing = await make_student(db, email="taken@campushire.dev")
    content = csv_bytes(
        [
            STUDENT_HEADER,
            ["good.one@campushire.dev", "Good One", "+919876543210", "CSE", "3.4", "E-1", "2027"],
            ["bad-email", "Bad Email", "+919876543211", "CSE", "3.0", "", ""],  # row 3
            ["gpa.bad@campushire.dev", "Gpa Bad", "+919876543212", "CSE", "4.5", "", ""],  # row 4
            ["good.two@campushire.dev", "Good Two", "", "ECE", "2.9", "", ""],
            [existing.email.upper(), "Dup Existing", "+919876543213", "CSE", "3.0", "", ""],  # row 6
            ["good.one@campushire.dev", "Dup In File", "+919876543214", "CSE", "3.0", "", ""],  # row 7
            ["phone.bad@campushire.dev", "Phone Bad", "12345", "CSE", "3.0", "", ""],  # row 8
            ["", "", "", "", "", "", ""],  # blank row skipped
        ]
    )
    resp = await _post_import(client, admin, "students", "students.csv", content)
    assert resp.status_code == 202, resp.text
    job = resp.json()
    assert job["type"] == "DATA_IMPORT" and job["params"]["entity"] == "students"
    assert ("data.import", {"job_id": job["id"]}) in captured.tasks

    await importer.run_import_job(job["id"], session_factory=session_factory, store=fake_storage)
    done = (await client.get(f"{API}/jobs/{job['id']}", headers=auth_headers(admin))).json()
    assert done["status"] == "SUCCEEDED" and done["progress"] == 100
    result = done["result"]
    assert result["rows_ok"] == 2 and result["rows_failed"] == 5 and result["total_rows"] == 7
    by_row = {e["row"]: e for e in result["errors"]}
    assert set(by_row) == {3, 4, 6, 7, 8}
    assert by_row[3]["field"] == "email" and by_row[4]["field"] == "gpa" and by_row[8]["field"] == "phone"
    assert "already exists" in by_row[6]["message"] and "row 2" in by_row[7]["message"]

    created = (
        (await db.execute(select(User).where(User.email.in_(["good.one@campushire.dev", "good.two@campushire.dev"]))))
        .scalars()
        .all()
    )
    assert len(created) == 2 and all(u.role == "STUDENT" and u.email_verified_at for u in created)
    students = (
        await db.execute(select(func.count()).select_from(Student).where(Student.user_id.in_([u.id for u in created])))
    ).scalar_one()
    assert students == 2
    resets = [e for e in captured.emails if e["template"] == "reset_password"]
    assert {e["to"] for e in resets} == {"good.one@campushire.dev", "good.two@campushire.dev"}
    assert (
        await db.execute(select(func.count()).select_from(AuditLog).where(AuditLog.action == "data.import"))
    ).scalar_one() == 1


async def test_import_companies_xlsx(client, db, admin, fake_storage, session_factory):
    await make_company(db, registration_number="U72200KA2015PTC082345")
    header = importer.COLUMNS["companies"]
    rows = [
        header,
        [
            "Acme Labs",
            "u74999mh2017ptc291234",
            "Pune",
            "Jane Doe",
            "hr@acme.example",
            "+919812345678",
            "Software",
            "https://acme.example",
            "Desc",
        ],
        ["Dupe Co", "U72200KA2015PTC082345", "Pune", "Jane Doe", "hr@dupe.example", "", "", "", ""],  # row 3 existing
        ["Bad Reg", "ABC", "Pune", "Jane Doe", "hr@bad.example", "", "", "", ""],  # row 4
        ["Intl Co", "US-DE5567123", "Remote", "Joe", "hr@intl.example", "", "", "", ""],
    ]
    resp = await _post_import(
        client, admin, "companies", "companies.xlsx", xlsx_bytes(rows), "application/octet-stream"
    )
    assert resp.status_code == 202, resp.text
    job = resp.json()
    await importer.run_import_job(job["id"], session_factory=session_factory, store=fake_storage)
    done = (await client.get(f"{API}/jobs/{job['id']}", headers=auth_headers(admin))).json()
    assert done["status"] == "SUCCEEDED"
    assert done["result"]["rows_ok"] == 2 and done["result"]["rows_failed"] == 2
    assert {e["row"] for e in done["result"]["errors"]} == {3, 4}
    acme = (
        await db.execute(select(Company).where(Company.registration_number == "U74999MH2017PTC291234"))
    ).scalar_one()
    assert acme.status == "ACTIVE" and acme.created_by == admin.id


async def test_import_numeric_xlsx_cells(client, db, admin, fake_storage, session_factory):
    rows = [STUDENT_HEADER, ["num@campushire.dev", "Numeric Cells", 919876543210, "CSE", 3.5, "N-1", 2027.0]]
    resp = await _post_import(client, admin, "students", "n.xlsx", xlsx_bytes(rows))
    job = resp.json()
    await importer.run_import_job(job["id"], session_factory=session_factory, store=fake_storage)
    got = (await client.get(f"{API}/jobs/{job['id']}", headers=auth_headers(admin))).json()
    assert got["result"]["rows_ok"] == 1, got["result"]


async def test_import_upload_validation(client, admin):
    ok_header = csv_bytes([STUDENT_HEADER])
    r = await _post_import(client, admin, "students", "x.txt", b"a,b")
    assert r.status_code == 422
    r = await _post_import(client, admin, "students", "empty.csv", b"")
    assert r.status_code == 422
    r = await _post_import(client, admin, "students", "big.csv", ok_header + b"a" * (5 * 1024 * 1024 + 1))
    assert r.status_code == 422 and r.json()["error"]["code"] == "FILE_TOO_LARGE"
    r = await _post_import(client, admin, "students", "fake.xlsx", b"not a zip file")
    assert r.status_code == 422
    r = await _post_import(
        client, admin, "students", "missing.csv", csv_bytes([["email", "full_name"], ["a@b.co", "A B"]])
    )
    assert r.status_code == 422 and "department" in r.json()["error"]["message"]
    r = await _post_import(client, admin, "internships", "x.csv", ok_header)
    assert r.status_code == 422


async def test_import_job_fails_on_bad_file(client, db, admin, fake_storage, session_factory):
    from app.modules.documents.models import Document

    doc = Document(
        id=uuid.uuid4(),
        owner_id=admin.id,
        kind="IMPORT",
        bucket="documents",
        object_key=f"{admin.id}/x/bad.csv",
        filename="bad.csv",
        content_type="text/csv",
        size_bytes=3,
        status="UPLOADED",
        verification_status="VERIFIED",
    )
    db.add(doc)
    job = Job(
        id=uuid.uuid4(),
        type="DATA_IMPORT",
        status="QUEUED",
        requested_by=admin.id,
        params={"entity": "students", "document_id": str(doc.id), "filename": "bad.csv"},
    )
    db.add(job)
    await db.flush()
    fake_storage.put(doc.object_key, b"foo,bar\n1,2\n", "text/csv")
    await importer.run_import_job(job.id, session_factory=session_factory, store=fake_storage)
    got = (await client.get(f"{API}/jobs/{job.id}", headers=auth_headers(admin))).json()
    assert got["status"] == "FAILED" and "Missing required column" in got["error"]


# ---- jobs ------------------------------------------------------------------------------------------------------------------------
async def test_jobs_list_filters_and_404(client, db, admin):
    for jt, st in [("DATA_EXPORT", "SUCCEEDED"), ("DATA_IMPORT", "FAILED"), ("REPORT_EXPORT", "SUCCEEDED")]:
        db.add(Job(type=jt, status=st, requested_by=admin.id, params={}))
    await db.flush()
    h = auth_headers(admin)
    body = (await client.get(f"{API}/jobs", headers=h)).json()
    assert body["total"] == 2  # REPORT_EXPORT jobs belong to /jobs (reports module)
    assert (await client.get(f"{API}/jobs?status=FAILED", headers=h)).json()["total"] == 1
    assert (await client.get(f"{API}/jobs?type=DATA_EXPORT", headers=h)).json()["total"] == 1
    assert (await client.get(f"{API}/jobs/{uuid.uuid4()}", headers=h)).status_code == 404
    assert (await client.get(f"{API}/jobs/{uuid.uuid4()}/download", headers=h)).status_code == 404
