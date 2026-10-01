"""Data export (CSV / XLSX) for students, companies, internships, applications and feedback.

``run_export_job(job_id)`` is called by the Celery task ``data.export`` through ``runtime.run_async``.
It builds the file in memory, stores it in the ``reports`` bucket, adds a ``documents`` row (kind EXPORT)
and marks the job SUCCEEDED. Cells that start with ``= + - @`` are prefixed with ``'`` (CSV/formula injection).
"""

import csv
import io
import logging
import re
import uuid
from collections.abc import Callable
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from sqlalchemy import Select, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import side_effects
from app.core.base import utcnow
from app.core.config import settings
from app.core.storage import Storage
from app.core.storage import storage as default_storage
from app.modules.admin import repository as repo
from app.modules.admin import schemas as s
from app.modules.admin.models import Job
from app.modules.applications.models import Application
from app.modules.companies.models import Company
from app.modules.documents.models import Document
from app.modules.feedback.models import StudentFeedback
from app.modules.internships.models import Internship
from app.modules.students.models import Student
from app.modules.users.models import User

log = logging.getLogger(__name__)

MAX_EXPORT_ROWS = 100_000
CSV_TYPE = "text/csv"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


# ---- filters ------------------------------------------------------------------------------------------------
def _str(v: Any) -> str:
    if not isinstance(v, str) or not v.strip():
        raise ValueError("must be a non-empty string")
    return v.strip()


def _enum(*allowed: str) -> Callable[[Any], str]:
    def parse(v: Any) -> str:
        v = _str(v).upper()
        if v not in allowed:
            raise ValueError(f"must be one of {', '.join(allowed)}")
        return v

    return parse


def _dec(v: Any) -> Decimal:
    try:
        return Decimal(str(v))
    except Exception as exc:  # noqa: BLE001
        raise ValueError("must be a number") from exc


def _bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if isinstance(v, str) and v.lower() in ("true", "false"):
        return v.lower() == "true"
    raise ValueError("must be true or false")


def _uuid(v: Any) -> uuid.UUID:
    try:
        return uuid.UUID(str(v))
    except ValueError as exc:
        raise ValueError("must be a UUID") from exc


def _date(v: Any) -> date:
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError as exc:
        raise ValueError("must be a date (YYYY-MM-DD)") from exc


_APP_STATUS = ("PENDING", "UNDER_REVIEW", "SHORTLISTED", "INTERVIEW", "ACCEPTED", "REJECTED", "WITHDRAWN")
_INTERNSHIP_STATUS = ("DRAFT", "PENDING_APPROVAL", "APPROVED", "REJECTED", "CLOSED")

FILTER_PARSERS: dict[str, dict[str, Callable[[Any], Any]]] = {
    "students": {"department": _str, "gpa_min": _dec, "gpa_max": _dec, "is_active": _bool},
    "companies": {"status": _enum("PENDING", "ACTIVE", "ARCHIVED"), "industry": _str},
    "internships": {"status": _enum(*_INTERNSHIP_STATUS), "domain": _str, "company_id": _uuid},
    "applications": {"status": _enum(*_APP_STATUS), "internship_id": _uuid, "from": _date, "to": _date},
    "feedback": {"company_id": _uuid, "internship_id": _uuid, "from": _date, "to": _date},
}


def parse_filters(entity: str, filters: dict[str, Any]) -> dict[str, Any]:
    """Validate + coerce ``filters``; raises ``ValueError`` with ``"<key>: <problem>"`` messages."""
    parsers = FILTER_PARSERS[entity]
    out: dict[str, Any] = {}
    for key, value in (filters or {}).items():
        if key not in parsers:
            raise ValueError(f"{key}: unknown filter for {entity} (allowed: {', '.join(sorted(parsers))})")
        if value is None or value == "":
            continue
        try:
            out[key] = parsers[key](value)
        except ValueError as exc:
            raise ValueError(f"{key}: {exc}") from exc
    return out


def _day_start(d: date) -> datetime:
    return datetime(d.year, d.month, d.day, tzinfo=UTC)


# ---- row sources ---------------------------------------------------------------------------------------------
Dataset = tuple[list[str], Select[Any]]


def _students(f: dict[str, Any]) -> Dataset:
    apps = (
        select(func.count()).select_from(Application).where(Application.student_id == Student.user_id).scalar_subquery()
    )
    stmt = (
        select(
            User.id,
            User.email,
            User.full_name,
            User.phone,
            Student.department,
            Student.gpa,
            Student.enrollment_no,
            Student.graduation_year,
            User.is_active,
            User.email_verified_at,
            apps,
            User.created_at,
        )
        .join(Student, Student.user_id == User.id)
        .order_by(User.created_at, User.id)
    )
    if "department" in f:
        stmt = stmt.where(Student.department == f["department"])
    if "gpa_min" in f:
        stmt = stmt.where(Student.gpa >= f["gpa_min"])
    if "gpa_max" in f:
        stmt = stmt.where(Student.gpa <= f["gpa_max"])
    if "is_active" in f:
        stmt = stmt.where(User.is_active.is_(f["is_active"]))
    cols = [
        "user_id",
        "email",
        "full_name",
        "phone",
        "department",
        "gpa",
        "enrollment_no",
        "graduation_year",
        "is_active",
        "email_verified_at",
        "application_count",
        "created_at",
    ]
    return cols, stmt


def _companies(f: dict[str, Any]) -> Dataset:
    interns = select(func.count()).select_from(Internship).where(Internship.company_id == Company.id).scalar_subquery()
    stmt = select(
        Company.id,
        Company.name,
        Company.registration_number,
        Company.industry,
        Company.location,
        Company.website,
        Company.contact_person_name,
        Company.contact_email,
        Company.contact_phone,
        Company.status,
        interns,
        Company.created_at,
    ).order_by(Company.created_at, Company.id)
    if "status" in f:
        stmt = stmt.where(Company.status == f["status"])
    if "industry" in f:
        stmt = stmt.where(Company.industry == f["industry"])
    cols = [
        "id",
        "name",
        "registration_number",
        "industry",
        "location",
        "website",
        "contact_person_name",
        "contact_email",
        "contact_phone",
        "status",
        "internship_count",
        "created_at",
    ]
    return cols, stmt


def _internships(f: dict[str, Any]) -> Dataset:
    apps = (
        select(func.count())
        .select_from(Application)
        .where(Application.internship_id == Internship.id)
        .scalar_subquery()
    )
    stmt = (
        select(
            Internship.id,
            Internship.title,
            Company.name,
            Internship.domain,
            Internship.location,
            Internship.work_mode,
            Internship.stipend_monthly,
            Internship.currency,
            Internship.duration_weeks,
            Internship.start_date,
            Internship.end_date,
            Internship.application_deadline,
            Internship.openings,
            Internship.status,
            apps,
            Internship.archived_at,
            Internship.created_at,
        )
        .join(Company, Company.id == Internship.company_id)
        .order_by(Internship.created_at, Internship.id)
    )
    if "status" in f:
        stmt = stmt.where(Internship.status == f["status"])
    if "domain" in f:
        stmt = stmt.where(Internship.domain == f["domain"])
    if "company_id" in f:
        stmt = stmt.where(Internship.company_id == f["company_id"])
    cols = [
        "id",
        "title",
        "company",
        "domain",
        "location",
        "work_mode",
        "stipend_monthly",
        "currency",
        "duration_weeks",
        "start_date",
        "end_date",
        "application_deadline",
        "openings",
        "status",
        "application_count",
        "archived_at",
        "created_at",
    ]
    return cols, stmt


def _applications(f: dict[str, Any]) -> Dataset:
    stmt = (
        select(
            Application.id,
            User.full_name,
            User.email,
            Student.department,
            Internship.title,
            Company.name,
            Application.status,
            Application.status_changed_at,
            Application.completed_at,
            Application.created_at,
        )
        .join(Student, Student.user_id == Application.student_id)
        .join(User, User.id == Student.user_id)
        .join(Internship, Internship.id == Application.internship_id)
        .join(Company, Company.id == Internship.company_id)
        .order_by(Application.created_at, Application.id)
    )
    if "status" in f:
        stmt = stmt.where(Application.status == f["status"])
    if "internship_id" in f:
        stmt = stmt.where(Application.internship_id == f["internship_id"])
    if "from" in f:
        stmt = stmt.where(Application.created_at >= _day_start(f["from"]))
    if "to" in f:
        stmt = stmt.where(Application.created_at < _day_start(f["to"]) + timedelta(days=1))
    cols = [
        "id",
        "student_name",
        "student_email",
        "department",
        "internship",
        "company",
        "status",
        "status_changed_at",
        "completed_at",
        "created_at",
    ]
    return cols, stmt


def _feedback(f: dict[str, Any]) -> Dataset:
    stmt = (
        select(
            StudentFeedback.id,
            User.full_name,
            Company.name,
            Internship.title,
            StudentFeedback.company_culture,
            StudentFeedback.mentorship,
            StudentFeedback.technical_learning,
            StudentFeedback.work_environment,
            StudentFeedback.overall,
            StudentFeedback.comments,
            StudentFeedback.suggestions,
            StudentFeedback.is_anonymous,
            StudentFeedback.response_body,
            StudentFeedback.created_at,
        )
        .join(User, User.id == StudentFeedback.student_id)
        .join(Company, Company.id == StudentFeedback.company_id)
        .join(Internship, Internship.id == StudentFeedback.internship_id)
        .order_by(StudentFeedback.created_at, StudentFeedback.id)
    )
    if "company_id" in f:
        stmt = stmt.where(StudentFeedback.company_id == f["company_id"])
    if "internship_id" in f:
        stmt = stmt.where(StudentFeedback.internship_id == f["internship_id"])
    if "from" in f:
        stmt = stmt.where(StudentFeedback.created_at >= _day_start(f["from"]))
    if "to" in f:
        stmt = stmt.where(StudentFeedback.created_at < _day_start(f["to"]) + timedelta(days=1))
    cols = [
        "id",
        "student_name",
        "company",
        "internship",
        "company_culture",
        "mentorship",
        "technical_learning",
        "work_environment",
        "overall",
        "comments",
        "suggestions",
        "is_anonymous",
        "response",
        "created_at",
    ]
    return cols, stmt


DATASETS: dict[str, Callable[[dict[str, Any]], Dataset]] = {
    "students": _students,
    "companies": _companies,
    "internships": _internships,
    "applications": _applications,
    "feedback": _feedback,
}


# ---- rendering -------------------------------------------------------------------------------------------------
def to_cell(value: Any) -> Any:
    """Normalise a DB value into a CSV/XLSX-safe cell (formula-injection guarded)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ") if value.tzinfo else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, uuid.UUID):
        return str(value)
    if isinstance(value, (list, tuple)):
        value = ", ".join(str(v) for v in value)
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def render_csv(columns: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO(newline="")
    w = csv.writer(buf)
    w.writerow(columns)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def render_xlsx(columns: list[str], rows: list[list[Any]], title: str = "export") -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=title[:31])
    ws.append(columns)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


async def build_export(session: AsyncSession, entity: str, fmt: str, filters: dict[str, Any]) -> tuple[bytes, int]:
    """Return ``(file_bytes, row_count)``."""
    columns, stmt = DATASETS[entity](parse_filters(entity, filters))
    result = await session.execute(stmt.limit(MAX_EXPORT_ROWS))
    rows = [[to_cell(v) for v in row] for row in result.all()]
    data = render_csv(columns, rows) if fmt == "csv" else render_xlsx(columns, rows, entity)
    return data, len(rows)


# ---- job runner --------------------------------------------------------------------------------------------------
_SLUG = re.compile(r"[^A-Za-z0-9._-]+")


async def _fail(session: AsyncSession, job_id: uuid.UUID, message: str) -> None:
    await session.rollback()
    side_effects.discard(session)
    job = await repo.get_job(session, job_id)
    if job is None:
        return
    job.status = "FAILED"
    job.error = message[:2000]
    job.finished_at = utcnow()
    await session.commit()
    await _publish(job)


async def _publish(job: Job, download_url: str | None = None) -> None:
    if job.requested_by is None:
        return
    try:
        await side_effects.dispatcher.publish(job.requested_by, {"type": "job", "data": _job_event(job, download_url)})
    except Exception:  # noqa: BLE001
        log.exception("could not publish job event %s", job.id)


def _job_event(job: Job, download_url: str | None = None) -> dict[str, Any]:
    return s.job_out(job, download_url).model_dump(mode="json")


async def run_export_job(
    job_id: uuid.UUID | str,
    *,
    session_factory: repo.SessionFactory | None = None,
    store: Storage | None = None,
) -> None:
    """Execute a DATA_EXPORT job end to end. Never raises: failures are recorded on the job."""
    job_id = uuid.UUID(str(job_id))
    store = store or default_storage
    async with repo.worker_session(session_factory) as session:
        job = await repo.get_job(session, job_id)
        if job is None:
            log.error("export job %s not found", job_id)
            return
        try:
            params = job.params or {}
            entity = str(params.get("entity", ""))
            fmt = str(params.get("format", "csv"))
            if entity not in DATASETS or fmt not in ("csv", "xlsx"):
                raise ValueError(f"Invalid export parameters: entity={entity!r} format={fmt!r}")
            if job.requested_by is None:
                raise ValueError("Job has no requester")
            job.status = "RUNNING"
            job.started_at = utcnow()
            job.progress = 10
            await session.commit()

            data, count = await build_export(session, entity, fmt, params.get("filters") or {})
            job = await repo.get_job(session, job_id)
            assert job is not None
            job.progress = 70
            await session.commit()

            doc_id = uuid.uuid4()
            stamp = utcnow().strftime("%Y%m%d-%H%M%S")
            filename = f"{entity}-export-{stamp}.{fmt}"
            content_type = CSV_TYPE if fmt == "csv" else XLSX_TYPE
            key = f"{job.requested_by}/{doc_id}/{_SLUG.sub('-', filename)}"
            await store.put_object(settings.BUCKET_REPORTS, key, data, content_type)

            session.add(
                Document(
                    id=doc_id,
                    owner_id=job.requested_by,
                    kind="EXPORT",
                    bucket=settings.BUCKET_REPORTS,
                    object_key=key,
                    filename=filename,
                    content_type=content_type,
                    size_bytes=max(len(data), 1),
                    status="UPLOADED",
                    verification_status="VERIFIED",
                )
            )
            await session.flush()
            job.result_document_id = doc_id
            job.result = {"rows_ok": count, "rows_failed": 0, "entity": entity, "format": fmt}
            job.status = "SUCCEEDED"
            job.progress = 100
            job.finished_at = utcnow()
            side_effects.notify(
                session,
                job.requested_by,
                "SYSTEM",
                "Export ready",
                f"Your {entity} export ({count} rows) is ready to download.",
                link="/admin/data",
                data={"job_id": str(job.id)},
            )
            side_effects.audit(
                session,
                job.requested_by,
                "data.export.complete",
                "job",
                job.id,
                after={"entity": entity, "rows": count},
            )
            await session.commit()
            await side_effects.flush(session)
            url = None
            try:
                url = await store.presign_get(settings.BUCKET_REPORTS, key, filename)
            except Exception:  # noqa: BLE001
                log.debug("presign failed", exc_info=True)
            await _publish(job, url)
        except Exception as exc:  # noqa: BLE001
            log.exception("export job %s failed", job_id)
            await _fail(session, job_id, str(exc) or type(exc).__name__)
