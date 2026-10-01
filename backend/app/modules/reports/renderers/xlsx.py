"""XLSX renderer (openpyxl): Summary sheet (title, params, KPIs), Charts sheet (chart data), one sheet per table."""

from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.reports.renderers.common import parse_iso, safe_text
from app.modules.reports.schemas import ReportData

_HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=16, color="312E81")
_THIN = Side(style="thin", color="D4D4D8")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_NUMBER_FORMATS = {
    "number": "#,##0.##",
    "percent": '0.0"%"',
    "currency": '"INR" #,##0',
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm",
}


def _sheet_title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", name).strip()[:28] or "Sheet"
    title, i = base, 2
    while title.lower() in used:
        title = f"{base[:25]} {i}"
        i += 1
    used.add(title.lower())
    return title


def _cell_value(value: Any, fmt: str | None) -> Any:
    if value is None:
        return None
    if fmt in ("date", "datetime"):
        dt = parse_iso(value)
        if dt is not None:
            return dt.replace(tzinfo=None)
        return safe_text(str(value))
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    return safe_text(str(value))


def _header(ws: Worksheet, row: int, labels: list[str]) -> None:
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font, c.border = _HEADER_FILL, _HEADER_FONT, _BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(ws: Worksheet, max_width: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            text = c.value.strftime("%Y-%m-%d %H:%M") if hasattr(c.value, "strftime") else str(c.value)
            widths[c.column] = max(widths.get(c.column, 8), min(len(text) + 2, max_width))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def render_xlsx(report: ReportData) -> bytes:
    wb = Workbook()
    used: set[str] = set()

    # ---- Summary
    ws = wb.active
    assert ws is not None
    ws.title = _sheet_title("Summary", used)
    ws["A1"] = safe_text(report.title)
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = "Generated at"
    ws["B2"] = report.generated_at.replace(tzinfo=None)
    ws["B2"].number_format = _NUMBER_FORMATS["datetime"]
    ws["A2"].font = Font(bold=True)
    r = 3
    for k, v in report.params.items():
        ws.cell(row=r, column=1, value=safe_text(str(k))).font = Font(bold=True)
        ws.cell(row=r, column=2, value=safe_text(str(v)))
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Key metrics").font = Font(bold=True, size=12)
    r += 1
    _header(ws, r, ["Metric", "Value", "Note"])
    for k in report.kpis:
        r += 1
        ws.cell(row=r, column=1, value=safe_text(k.label)).border = _BORDER
        vc = ws.cell(row=r, column=2, value=_cell_value(k.value, None))
        vc.border = _BORDER
        if isinstance(k.value, (int, float)) and k.format in _NUMBER_FORMATS:
            vc.number_format = _NUMBER_FORMATS[k.format]
        ws.cell(row=r, column=3, value=safe_text(k.hint) if k.hint else None).border = _BORDER
    _autosize(ws)

    # ---- Chart data
    if report.charts:
        cs = wb.create_sheet(_sheet_title("Charts", used))
        row = 1
        for ch in report.charts:
            cs.cell(row=row, column=1, value=safe_text(ch.title)).font = Font(bold=True, size=12)
            row += 1
            if ch.type in ("line", "bar", "area"):
                _header(cs, row, [""] + [s.name for s in ch.series])
                for i, x in enumerate(ch.x):
                    row += 1
                    cs.cell(row=row, column=1, value=safe_text(x))
                    for j, s in enumerate(ch.series, start=2):
                        cs.cell(row=row, column=j, value=s.data[i] if i < len(s.data) else None)
            elif ch.type in ("pie", "donut"):
                _header(cs, row, ["Name", "Value"])
                for d in ch.data:
                    row += 1
                    cs.cell(row=row, column=1, value=safe_text(d.name))
                    cs.cell(row=row, column=2, value=d.value)
            else:  # radar
                _header(cs, row, [""] + [s.name for s in ch.series])
                for i, ind in enumerate(ch.indicators):
                    row += 1
                    cs.cell(row=row, column=1, value=safe_text(f"{ind.name} (max {ind.max})"))
                    for j, s in enumerate(ch.series, start=2):
                        cs.cell(row=row, column=j, value=s.data[i] if i < len(s.data) else None)
            row += 2
        _autosize(cs)

    # ---- One sheet per table
    for t in report.tables:
        ts = wb.create_sheet(_sheet_title(t.title, used))
        _header(ts, 1, [c.label for c in t.columns])
        for ri, rowdata in enumerate(t.rows, start=2):
            for ci, col in enumerate(t.columns, start=1):
                c = ts.cell(row=ri, column=ci, value=_cell_value(rowdata.get(col.key), col.format))
                c.border = _BORDER
                if col.format in _NUMBER_FORMATS and c.value is not None and not isinstance(c.value, str):
                    c.number_format = _NUMBER_FORMATS[col.format]
        ts.freeze_panes = "A2"
        if t.rows:
            ts.auto_filter.ref = f"A1:{get_column_letter(len(t.columns))}{len(t.rows) + 1}"
        _autosize(ts)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
